from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('AI Assisted Project Induction Platform - Context Snapshot', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Date: ' + __import__('datetime').datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
doc.add_paragraph()

# Project Info
doc.add_heading('Project Info', level=1)
info = [
    ['Location', r'C:\Users\K Abdul Waseem\company_project'],
    ['Stack', 'React 19 + JavaScript, Vite, Tailwind CSS v3, react-router-dom, axios, react-hook-form, lucide-react'],
    ['Mode', 'Frontend-only with mock provider (VITE_API_MODE=mock)'],
    ['Design', 'Microsoft Copilot/Learn/Azure inspired (light/dark theme)'],
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Field'
hdr[1].text = 'Value'
for field, value in info:
    row = table.add_row().cells
    row[0].text = field
    row[1].text = value
doc.add_paragraph()

# Verification Status
doc.add_heading('Verification Status', level=1)
ver = [
    ['Lint', '0 errors (7 warnings)'],
    ['Build', 'Passed - dist/ generated with all assets'],
    ['Routes', 'All routes, guards (ProtectedRoute/RoleRoute/PublicOnlyRoute), and status pages wired'],
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Check'
hdr[1].text = 'Status'
for check, status in ver:
    row = table.add_row().cells
    row[0].text = check
    row[1].text = status
doc.add_paragraph()

# Completed Work
doc.add_heading('Completed Work', level=1)
work_items = [
    ('Preferences System', 'PreferencesContext + usePreferences hook with persisted profile/notification settings'),
    ('Settings Page', 'Appearance (light/dark), Profile (display name/contact/timezone), Workspace (notifications, reset demo data, sign out)'),
    ('Quiz Flow', 'Rewritten QuizAttemptPage to match service contract (get(), getQuestions(), submit() with flat answers array). Extended QuizTimer (onTimeChange), QuestionNavigator (flagged), QuizQuestion (flag support)'),
    ('Notifications', 'Client-side filtering by user preferences, markRead/markAllRead return filtered arrays'),
    ('Provider Wiring', 'PreferencesProvider added to App.jsx, effective user propagated to Sidebar/UserMenu/Dashboard'),
    ('Lint Fixes', 'Resolved setState-in-effect, ref access during render, and React Hooks issues'),
]
for title_item, desc in work_items:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(title_item + ': ').bold = True
    p.add_run(desc)
doc.add_paragraph()

# Key Files
doc.add_heading('Key Files', level=1)
files = [
    'src/App.jsx', 'src/routes/AppRoutes.jsx',
    'src/contexts/PreferencesContext.jsx', 'src/hooks/usePreferences.js',
    'src/pages/SettingsPage.jsx',
    'src/pages/learner/QuizAttemptPage.jsx',
    'src/components/quiz/QuizTimer.jsx', 'src/components/quiz/QuestionNavigator.jsx', 'src/components/quiz/QuizQuestion.jsx',
    'src/services/notifications/notificationService.js',
    'src/constants/storage.js', 'src/utils/date.js',
]
for f in files:
    doc.add_paragraph(f, style='List Bullet')
doc.add_paragraph()

# Commands
doc.add_heading('Commands', level=1)
commands = [
    ['npm run dev', 'Start dev server'],
    ['npm run lint', 'Run ESLint'],
    ['npm run build', 'Production build'],
    ['npm run preview', 'Preview built app'],
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Command'
hdr[1].text = 'Purpose'
for cmd, purpose in commands:
    row = table.add_row().cells
    row[0].text = cmd
    row[1].text = purpose

doc.save('AI_Assisted_Project_Induction_Platform_Context.docx')
print('Created: AI_Assisted_Project_Induction_Platform_Context.docx')
